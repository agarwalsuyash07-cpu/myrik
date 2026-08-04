#!/usr/bin/env python3
"""
split_orders_by_time.py

Reads an order-export .xlsx file and splits orders into two groups:
  - "Before 5PM": orders created before 17:00, sorted by location
  - "After 5PM":  orders created at/after 17:00, sorted by location

Each row has 3 columns: Customer Number, Location, Order Creation Date (dd/mm/yyyy).

Usage:
    python3 split_orders_by_time.py <input.xlsx>              # prints both groups
    python3 split_orders_by_time.py <input.xlsx> output.xlsx  # writes a 2-sheet workbook

With no output path the rows are printed as tab-separated text, so they paste
straight into a spreadsheet. Pass an output path to get the workbook instead.

Notes:
- Works with export files where the real header row (e.g. "Customer Phone",
  "City", "Created At") isn't row 1 — it scans the first 20 rows to find it.
- Change CUTOFF_HOUR below to split at a different time.
- Change PHONE_COL / CITY_COL / DATE_COL header names below if your export
  uses different column names.
"""

import sys
from pathlib import Path
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# ---- configuration -------------------------------------------------------
CUTOFF_HOUR = 17          # 5 PM cutoff; orders before this hour -> "Before" sheet
PHONE_COL = "Customer Phone"
CITY_COL = "City"
DATE_COL = "Created At"
DATE_INPUT_FORMAT = "%d-%m-%Y %H:%M:%S"   # format used in the source file
DATE_OUTPUT_FORMAT = "%d/%m/%Y"           # format required in the output
HEADER_SCAN_ROWS = 20      # how many rows to scan looking for the header row
# ---------------------------------------------------------------------------


def find_header_row(ws, required_cols, max_scan=HEADER_SCAN_ROWS):
    """Return (row_index, {col_name: col_index}) for the first row containing
    all required_cols. Row/col indices are 1-based."""
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        values = {cell.value: cell.column for cell in row if cell.value is not None}
        if all(col in values for col in required_cols):
            return row[0].row, values
    raise ValueError(
        f"Could not find a header row containing {required_cols} "
        f"in the first {max_scan} rows."
    )


def parse_datetime(value):
    """Accept either a datetime object (already parsed by openpyxl) or a
    string in DATE_INPUT_FORMAT."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        return datetime.strptime(value.strip(), DATE_INPUT_FORMAT)
    raise ValueError(f"Unrecognized date value: {value!r}")


def load_orders(input_path):
    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.active

    header_row, col_map = find_header_row(ws, [PHONE_COL, CITY_COL, DATE_COL])
    phone_idx = col_map[PHONE_COL]
    city_idx = col_map[CITY_COL]
    date_idx = col_map[DATE_COL]

    before, after = [], []
    for row in ws.iter_rows(min_row=header_row + 1):
        phone = row[phone_idx - 1].value
        city = row[city_idx - 1].value
        raw_date = row[date_idx - 1].value

        if phone is None and city is None and raw_date is None:
            continue  # skip blank rows

        dt = parse_datetime(raw_date)
        entry = (phone, city, dt)

        if dt.hour < CUTOFF_HOUR:
            before.append(entry)
        else:
            after.append(entry)

    return before, after


COLUMNS = ["Customer Number", "Location", "Order Creation Date"]


def sort_orders(data):
    """Sort by location, then by date/time within each location."""
    return sorted(data, key=lambda x: (str(x[1]), x[2]))


def print_split(before, after):
    """Print both groups as tab-separated rows, ready to paste into a sheet.
    Groups are separated by a form feed; the web frontend renders one box per group."""
    for i, (title, rows) in enumerate((("Before 5PM", before), ("After 5PM", after))):
        if i:
            print("\f", end="")
        print(f"{title}: {len(rows)} orders")
        print("\t".join(COLUMNS))
        for phone, city, dt in sort_orders(rows):
            print(f"{phone}\t{city}\t{dt.strftime(DATE_OUTPUT_FORMAT)}")


def write_sheet(ws, data):
    headers = COLUMNS

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    body_font = Font(name="Arial")

    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for r, (phone, city, dt) in enumerate(sort_orders(data), start=2):
        ws.cell(row=r, column=1, value=str(phone)).font = body_font
        ws.cell(row=r, column=2, value=city).font = body_font
        ws.cell(row=r, column=3, value=dt.strftime(DATE_OUTPUT_FORMAT)).font = body_font

    for col, width in zip("ABC", [18, 20, 20]):
        ws.column_dimensions[col].width = width


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 split_orders_by_time.py <input.xlsx> [output.xlsx]")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    if not input_path.exists():
        print(f"Input file not found: {input_path}")
        sys.exit(1)

    before, after = load_orders(input_path)

    if len(sys.argv) < 3:
        print_split(before, after)
        return

    output_path = Path(sys.argv[2])
    wb_out = openpyxl.Workbook()
    ws_before = wb_out.active
    ws_before.title = "Before 5PM"
    ws_after = wb_out.create_sheet("After 5PM")

    write_sheet(ws_before, before)
    write_sheet(ws_after, after)

    wb_out.save(output_path)

    print(f"Before 5PM: {len(before)} orders")
    print(f"After 5PM:  {len(after)} orders")
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
